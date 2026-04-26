"""Spreadsheet backend: formula evaluation, cell dependency tracking, and data I/O."""

from __future__ import annotations

import csv
import re
from typing import Any, Callable

# ── Cell-address helpers ───────────────────────────────────────────────────────


def col_letter(col: int) -> str:
    """Zero-based column index → Excel-style letter(s).  0 → 'A', 26 → 'AA'."""
    s, n = "", col
    while True:
        s = chr(ord("A") + n % 26) + s
        n = n // 26 - 1
        if n < 0:
            break
    return s


def col_index(letters: str) -> int:
    """Excel-style column letters → zero-based index.  'A' → 0, 'AA' → 26."""
    n = 0
    for ch in letters.upper():
        n = n * 26 + (ord(ch) - ord("A") + 1)
    return n - 1


def a1_to_rc(ref: str) -> tuple[int, int]:
    """Parse an A1-style reference to a (row, col) zero-based pair."""
    m = re.fullmatch(r"([A-Za-z]+)(\d+)", ref.strip())
    if not m:
        raise ValueError(f"Bad cell reference: {ref}")
    letters, digits = m.groups()
    return int(digits) - 1, col_index(letters)


def rc_to_a1(row: int, col: int) -> str:
    """Convert a (row, col) zero-based pair to an A1-style string."""
    return f"{col_letter(col)}{row + 1}"


# ── Formula evaluation ─────────────────────────────────────────────────────────

CellKey = tuple[int, int]


class FormulaError(Exception):
    """Raised when a formula cannot be evaluated."""


def _product(xs: list[float]) -> float:
    p = 1.0
    for x in xs:
        p *= x
    return p


FUNCTIONS: dict[str, Callable[[list[float]], Any]] = {
    "SUM":     lambda xs: sum(xs),
    "AVERAGE": lambda xs: (sum(xs) / len(xs)) if xs else FormulaError("#DIV/0!"),
    "MIN":     lambda xs: min(xs) if xs else FormulaError("#VALUE!"),
    "MAX":     lambda xs: max(xs) if xs else FormulaError("#VALUE!"),
    "COUNT":   lambda xs: float(len(xs)),
    "PRODUCT": _product,
    "ABS":     lambda xs: abs(xs[0]) if len(xs) == 1 else FormulaError("#VALUE!"),
    "ROUND":   lambda xs: round(xs[0], int(xs[1])) if len(xs) == 2 else FormulaError("#VALUE!"),
}


class Evaluator:
    """Recursive-descent parser and evaluator for spreadsheet formula strings."""

    def __init__(self, get_cell: Callable[[int, int], Any]) -> None:
        self.get_cell = get_cell
        self.deps: set[CellKey] = set()
        self.tokens: list[tuple[str, Any]] = []
        self.pos: int = 0

    def evaluate(self, text: str) -> Any:
        """Parse and evaluate *text* (the formula body, without the leading '=')."""
        self.deps = set()
        self.tokens = self._tokenize(text)
        self.pos = 0
        result = self._expr()
        if self.pos != len(self.tokens):
            raise FormulaError("#SYNTAX")
        return result

    # ── Tokeniser ──────────────────────────────────────────────────────────

    def _tokenize(self, s: str) -> list[tuple[str, Any]]:
        tokens: list[tuple[str, Any]] = []
        i = 0
        while i < len(s):
            c = s[i]
            if c.isspace():
                i += 1
            elif c.isdigit() or (c == "." and i + 1 < len(s) and s[i + 1].isdigit()):
                j = i
                while j < len(s) and (s[j].isdigit() or s[j] == "."):
                    j += 1
                tokens.append(("NUM", float(s[i:j])))
                i = j
            elif c == '"':
                j = i + 1
                while j < len(s) and s[j] != '"':
                    j += 1
                tokens.append(("STR", s[i + 1 : j]))
                i = j + 1
            elif c.isalpha():
                j = i
                while j < len(s) and (s[j].isalnum() or s[j] == "_"):
                    j += 1
                w = s[i:j]
                if re.fullmatch(r"[A-Za-z]+\d+", w):
                    tokens.append(("REF", w.upper()))
                elif w.upper() in ("TRUE", "FALSE"):
                    tokens.append(("BOOL", w.upper() == "TRUE"))
                else:
                    tokens.append(("FUNC", w.upper()))
                i = j
            elif c in "+-*/^(),:<>=":
                if c in "<>" and i + 1 < len(s) and s[i + 1] == "=":
                    tokens.append(("OP", c + "="))
                    i += 2
                elif c == "<" and i + 1 < len(s) and s[i + 1] == ">":
                    tokens.append(("OP", "<>"))
                    i += 2
                else:
                    tokens.append(("OP", c))
                    i += 1
            else:
                raise FormulaError(f"#SYNTAX: unexpected '{c}'")
        return tokens

    # ── Parser helpers ─────────────────────────────────────────────────────

    def _peek(self, offset: int = 0) -> tuple[str | None, Any]:
        idx = self.pos + offset
        return self.tokens[idx] if idx < len(self.tokens) else (None, None)

    def _eat(self) -> tuple[str, Any]:
        t = self.tokens[self.pos]
        self.pos += 1
        return t

    def _match(self, kind: str, val: Any = None) -> tuple[str, Any] | None:
        t = self._peek()
        if t[0] == kind and (val is None or t[1] == val):
            return self._eat()
        return None

    # ── Grammar (precedence climbing) ──────────────────────────────────────

    def _expr(self) -> Any:
        left = self._additive()
        while self._peek()[0] == "OP" and self._peek()[1] in ("=", "<>", "<", ">", "<=", ">="):
            op = self._eat()[1]
            right = self._additive()
            left = {"=": left == right, "<>": left != right, "<": left < right,
                    ">": left > right, "<=": left <= right, ">=": left >= right}[op]
        return left

    def _additive(self) -> Any:
        left = self._term()
        while self._peek()[0] == "OP" and self._peek()[1] in ("+", "-"):
            op = self._eat()[1]
            right = self._term()
            left = (self._num(left) + self._num(right)) if op == "+" else (self._num(left) - self._num(right))
        return left

    def _term(self) -> Any:
        left = self._power()
        while self._peek()[0] == "OP" and self._peek()[1] in ("*", "/"):
            op = self._eat()[1]
            right = self._power()
            if op == "*":
                left = self._num(left) * self._num(right)
            else:
                r = self._num(right)
                if r == 0:
                    raise FormulaError("#DIV/0!")
                left = self._num(left) / r
        return left

    def _power(self) -> Any:
        left = self._factor()
        if self._peek()[0] == "OP" and self._peek()[1] == "^":
            self._eat()
            return self._num(left) ** self._num(self._factor())
        return left

    def _factor(self) -> Any:
        t = self._peek()
        if t[0] == "OP" and t[1] == "-":
            self._eat()
            return -self._num(self._factor())
        if t[0] == "OP" and t[1] == "+":
            self._eat()
            return self._num(self._factor())
        if t[0] == "NUM":
            return self._eat()[1]
        if t[0] == "STR":
            return self._eat()[1]
        if t[0] == "BOOL":
            return self._eat()[1]
        if t[0] == "OP" and t[1] == "(":
            self._eat()
            v = self._expr()
            if not self._match("OP", ")"):
                raise FormulaError("#SYNTAX: missing )")
            return v
        if t[0] == "REF":
            r1 = self._eat()[1]
            if self._peek()[0] == "OP" and self._peek()[1] == ":":
                self._eat()
                if self._peek()[0] != "REF":
                    raise FormulaError("#SYNTAX")
                r2 = self._eat()[1]
                return self._resolve_range(r1, r2)
            return self._resolve_cell(r1)
        if t[0] == "FUNC":
            name = self._eat()[1]
            if not self._match("OP", "("):
                raise FormulaError(f"#NAME?: {name}")
            args: list[Any] = []
            if not (self._peek()[0] == "OP" and self._peek()[1] == ")"):
                args.append(self._expr())
                while self._peek()[0] == "OP" and self._peek()[1] == ",":
                    self._eat()
                    args.append(self._expr())
            if not self._match("OP", ")"):
                raise FormulaError(f"#SYNTAX: missing ) in {name}")
            return self._call(name, args)
        raise FormulaError("#SYNTAX: unexpected token")

    # ── Reference resolution ───────────────────────────────────────────────

    def _resolve_cell(self, ref: str) -> Any:
        r, c = a1_to_rc(ref)
        self.deps.add((r, c))
        v = self.get_cell(r, c)
        if v is None or v == "":
            return 0
        if isinstance(v, str):
            try:
                return float(v)
            except ValueError:
                return v
        return v

    def _resolve_range(self, r1: str, r2: str) -> list[float]:
        rr1, cc1 = a1_to_rc(r1)
        rr2, cc2 = a1_to_rc(r2)
        rs, re_ = min(rr1, rr2), max(rr1, rr2)
        cs, ce = min(cc1, cc2), max(cc1, cc2)
        vals: list[float] = []
        for r in range(rs, re_ + 1):
            for c in range(cs, ce + 1):
                self.deps.add((r, c))
                v = self.get_cell(r, c)
                if v is None or v == "":
                    continue
                if isinstance(v, (int, float)):
                    vals.append(float(v))
                elif isinstance(v, str):
                    try:
                        vals.append(float(v))
                    except ValueError:
                        pass
        return vals

    def _call(self, name: str, args: list[Any]) -> Any:
        if name == "IF":
            if len(args) != 3:
                raise FormulaError("#VALUE!")
            return args[1] if args[0] else args[2]
        if name in ("CONCAT", "CONCATENATE"):
            return "".join(str(a) for a in args)
        if name not in FUNCTIONS:
            raise FormulaError(f"#NAME?: {name}")
        flat: list[float] = []
        for a in args:
            if isinstance(a, list):
                flat.extend(a)
            elif isinstance(a, (int, float)):
                flat.append(float(a))
            elif isinstance(a, str):
                try:
                    flat.append(float(a))
                except ValueError:
                    pass
            elif isinstance(a, bool):
                flat.append(1.0 if a else 0.0)
        result = FUNCTIONS[name](flat)
        if isinstance(result, FormulaError):
            raise result
        return result

    def _num(self, v: Any) -> float:
        if isinstance(v, bool):
            return 1.0 if v else 0.0
        if isinstance(v, (int, float)):
            return float(v)
        if isinstance(v, str):
            try:
                return float(v)
            except ValueError:
                raise FormulaError("#VALUE!")
        if isinstance(v, list):
            raise FormulaError("#VALUE!")
        raise FormulaError("#VALUE!")


# ── Spreadsheet model ──────────────────────────────────────────────────────────


class Spreadsheet:
    """
    In-memory spreadsheet with formula evaluation, dependency tracking,
    and undo/redo support.
    """

    def __init__(self, rows: int = 20, cols: int = 10) -> None:
        self.rows = rows
        self.cols = cols
        self.raw: dict[CellKey, str] = {}
        self.values: dict[CellKey, Any] = {}
        self.deps_of: dict[CellKey, set[CellKey]] = {}
        self.dependents: dict[CellKey, set[CellKey]] = {}
        self._history: list[dict] = []
        self._future: list[dict] = []
        self._history_limit = 50

    # ── Snapshot / restore ─────────────────────────────────────────────────

    def _snapshot(self) -> dict:
        return {
            "raw": dict(self.raw),
            "values": dict(self.values),
            "deps_of": {k: set(v) for k, v in self.deps_of.items()},
            "dependents": {k: set(v) for k, v in self.dependents.items()},
            "rows": self.rows,
            "cols": self.cols,
        }

    def _restore(self, s: dict) -> None:
        self.raw = dict(s["raw"])
        self.values = dict(s["values"])
        self.deps_of = {k: set(v) for k, v in s["deps_of"].items()}
        self.dependents = {k: set(v) for k, v in s["dependents"].items()}
        self.rows, self.cols = s["rows"], s["cols"]

    def _push(self) -> None:
        self._history.append(self._snapshot())
        if len(self._history) > self._history_limit:
            self._history.pop(0)
        self._future.clear()

    def undo(self) -> bool:
        """Revert to the previous state. Returns True if successful."""
        if not self._history:
            return False
        self._future.append(self._snapshot())
        self._restore(self._history.pop())
        return True

    def redo(self) -> bool:
        """Re-apply a previously undone state. Returns True if successful."""
        if not self._future:
            return False
        self._history.append(self._snapshot())
        self._restore(self._future.pop())
        return True

    # ── Cell mutation ──────────────────────────────────────────────────────

    def set_cell(self, row: int, col: int, text: str, record_history: bool = True) -> None:
        """Write raw content to a cell and propagate recalculation."""
        if record_history:
            self._push()
        key: CellKey = (row, col)
        if key in self.deps_of:
            for d in self.deps_of[key]:
                self.dependents.get(d, set()).discard(key)
            del self.deps_of[key]
        if text is None or text == "":
            self.raw.pop(key, None)
            self.values.pop(key, None)
        else:
            self.raw[key] = text
        self._recalc(row, col)
        self._recalc_deps(key)

    def _recalc(self, row: int, col: int) -> None:
        key: CellKey = (row, col)
        text = self.raw.get(key, "")
        if text == "":
            self.values.pop(key, None)
            return
        if isinstance(text, str) and text.startswith("="):
            ev = Evaluator(lambda r, c: self.values.get((r, c)))
            try:
                result = ev.evaluate(text[1:])
                if isinstance(result, list):
                    result = result[0] if result else ""
                self.values[key] = result
                self.deps_of[key] = ev.deps
                for d in ev.deps:
                    self.dependents.setdefault(d, set()).add(key)
            except FormulaError as e:
                self.values[key] = str(e)
            except Exception:
                self.values[key] = "#ERROR!"
        else:
            try:
                self.values[key] = (
                    float(text) if ("." in text or "e" in text.lower()) else int(text)
                )
            except (ValueError, TypeError):
                self.values[key] = text

    def _recalc_deps(self, key: CellKey) -> None:
        seen: set[CellKey] = set()
        queue = list(self.dependents.get(key, set()))
        while queue:
            c = queue.pop(0)
            if c in seen:
                continue
            seen.add(c)
            self._recalc(*c)
            queue.extend(self.dependents.get(c, set()))

    def recalculate_all(self) -> None:
        """Multi-pass recalculation to resolve inter-cell dependencies."""
        for _ in range(5):
            changed = False
            for k in list(self.raw.keys()):
                before = self.values.get(k)
                self._recalc(*k)
                if self.values.get(k) != before:
                    changed = True
            if not changed:
                break

    # ── Cell read accessors ────────────────────────────────────────────────

    def get_raw(self, r: int, c: int) -> str:
        return self.raw.get((r, c), "")

    def get_value(self, r: int, c: int) -> Any:
        return self.values.get((r, c), "")

    def get_display(self, r: int, c: int) -> str:
        """Return a display-ready string for the computed cell value."""
        v = self.values.get((r, c), "")
        if isinstance(v, float):
            return str(int(v)) if v.is_integer() else f"{v:g}"
        return str(v)

    # ── Structural mutations ───────────────────────────────────────────────

    def insert_row(self, at: int) -> None:
        self._push()
        self.raw = {((r + 1 if r >= at else r), c): v for (r, c), v in self.raw.items()}
        self.rows += 1
        self._rebuild()

    def delete_row(self, at: int) -> None:
        self._push()
        self.raw = {((r - 1 if r > at else r), c): v for (r, c), v in self.raw.items() if r != at}
        self.rows = max(1, self.rows - 1)
        self._rebuild()

    def insert_col(self, at: int) -> None:
        self._push()
        self.raw = {(r, (c + 1 if c >= at else c)): v for (r, c), v in self.raw.items()}
        self.cols += 1
        self._rebuild()

    def delete_col(self, at: int) -> None:
        self._push()
        self.raw = {(r, (c - 1 if c > at else c)): v for (r, c), v in self.raw.items() if c != at}
        self.cols = max(1, self.cols - 1)
        self._rebuild()

    def _rebuild(self) -> None:
        self.values.clear()
        self.deps_of.clear()
        self.dependents.clear()
        for k in list(self.raw.keys()):
            self._recalc(*k)
        self.recalculate_all()

    # ── Bulk load ──────────────────────────────────────────────────────────

    def load_from_2d(self, data: list[list[Any]]) -> None:
        """Replace all content with a 2-D list of raw values."""
        self.raw.clear()
        self.values.clear()
        self.deps_of.clear()
        self.dependents.clear()
        if data:
            self.rows = max(self.rows, len(data))
            self.cols = max(self.cols, max(len(r) for r in data))
        for r, row in enumerate(data):
            for c, val in enumerate(row):
                if val is None or val == "":
                    continue
                self.set_cell(r, c, str(val), record_history=False)
        self.recalculate_all()
        self._history.clear()
        self._future.clear()

    # ── CSV / XLSX I/O ─────────────────────────────────────────────────────

    def save_csv(self, path: str) -> None:
        data = [["" for _ in range(self.cols)] for _ in range(self.rows)]
        for (r, c), v in self.values.items():
            if r < self.rows and c < self.cols:
                data[r][c] = v
        with open(path, "w", newline="") as f:
            csv.writer(f).writerows(data)

    def load_csv(self, path: str) -> None:
        with open(path, newline="") as f:
            self.load_from_2d(list(csv.reader(f)))

    def save_xlsx(self, path: str) -> None:
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        for (r, c), text in self.raw.items():
            ws.cell(row=r + 1, column=c + 1, value=text)
        wb.save(path)

    def load_xlsx(self, path: str) -> None:
        from openpyxl import load_workbook
        wb = load_workbook(path)
        ws = wb.active
        self.load_from_2d([list(row) for row in ws.iter_rows(values_only=True)])
